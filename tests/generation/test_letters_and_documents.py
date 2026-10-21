"""Letters (decision 84) and reading the person's documents (decision 86)."""

import io
import json

import pytest

from qanoun.generation.documents import DocumentError, read_document
from qanoun.generation.letter_docx import letter_docx
from qanoun.generation.letters import write_letter

from ..retrieval.test_search import ARTICLES
from .test_answer import ScriptedLLM

LEAVE = ARTICLES[1]  # Labour Code art. 231


class FileLLM(ScriptedLLM):
    """Records the files sent with the prompt."""

    def generate(self, system, user, *, json_schema=None, attachments=()):
        self.files = list(attachments)
        return super().generate(system, user, json_schema=json_schema)


def letter(body, used=(1,)):
    return {
        "subject": "طلب عطلة",
        "recipient": "[اسم المشغل]",
        "body": body,
        "used_sources": list(used),
    }


def test_a_letter_names_its_articles_and_leaves_placeholders():
    llm = ScriptedLLM(letter("طبقا للمادة 231 من مدونة الشغل، أطلب عطلتي السنوية. [الاسم الكامل]"))
    out = write_letter(llm, "request", "ar", "بغيت الكونجي ديالي", [LEAVE], ["أنا: خدام"])
    assert out.verified and out.cited == [LEAVE.id] and "[الاسم الكامل]" in out.body
    system, prompt = llm.prompts[0]
    assert "formal Modern Standard Arabic" in system and "WHAT THEY SELECTED" in prompt


def test_a_letter_citing_an_article_it_was_not_given_is_corrected():
    bad = letter("طبقا للمادة 999 من مدونة الشغل…")
    good = letter("طبقا للمادة 231 من مدونة الشغل…")
    llm = ScriptedLLM(bad, good)
    out = write_letter(llm, "notice", "ar", "story", [LEAVE])
    assert out.verified and out.attempts == 2 and "999" in llm.prompts[1][1]


def test_unknown_kinds_are_refused():
    with pytest.raises(ValueError):
        write_letter(ScriptedLLM(), "poem", "ar", "x", [LEAVE])


def test_the_word_file_is_right_to_left_in_arabic():
    from docx import Document

    content = letter_docx("طلب", "[المشغل]", "السطر الأول\nالسطر الثاني", "ar")
    doc = Document(io.BytesIO(content))
    texts = [p.text for p in doc.paragraphs]
    assert "الموضوع: طلب" in texts and "السطر الثاني" in texts
    assert all(
        p._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi")
        is not None
        for p in doc.paragraphs
    )  # noqa: E501
    french = Document(io.BytesIO(letter_docx("Objet", "", "Bonjour", "fr")))
    assert "Objet: Objet" in [p.text for p in french.paragraphs]


READING = {
    "kind": "contract", "summary": "عقد كراء.", "story": "كنكري شقة، نمرة مول الدار 0612345678",
    "theme": "home", "dates": [{"what": "بداية العقد", "date": "2024-01-01"}],
    "amounts": [{"what": "الكرا", "amount": "2000 درهم"}],
    "clauses": [
        {"quote": "يؤدى الكراء في أول كل شهر", "meaning": "خاصك تخلص فالأول ديال الشهر"},
        {"quote": "يتنازل المكتري عن كل تعويض", "meaning": "invented clause"},
    ],
}  # fmt: skip


def pdf_with(text: str) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((40, 60), text, fontsize=11)
    return doc.tobytes()


def test_a_text_pdf_is_read_here_and_only_real_clauses_are_kept(monkeypatch):
    import qanoun.generation.documents as documents

    text = "Contrat de bail. " * 20 + "يؤدى الكراء في أول كل شهر"
    monkeypatch.setattr(documents, "pdf_text", lambda data: (text, 1))
    llm = FileLLM(READING)
    out = read_document(llm, "application/pdf", b"%PDF-fake", "darija")
    assert llm.files == [] and "DOCUMENT\nContrat de bail." in llm.prompts[0][1]  # text, no file
    assert out.read_as == "text" and [c["quote"] for c in out.clauses] == [
        "يؤدى الكراء في أول كل شهر"
    ]
    assert out.clauses[0]["checked"] and "0612345678" not in out.story and out.theme == "home"


def test_a_photo_is_sent_as_an_image_and_its_clauses_are_marked_unchecked():
    llm = FileLLM(READING)
    out = read_document(llm, "image/jpeg", b"\xff\xd8 photo", "ar")
    assert llm.files == [("image/jpeg", b"\xff\xd8 photo")] and out.read_as == "image"
    assert len(out.clauses) == 2 and not any(c["checked"] for c in out.clauses)


def test_a_real_pdf_without_text_is_read_as_an_image():
    llm = FileLLM(READING)
    data = pdf_with("x")  # almost no text: a scan
    out = read_document(llm, "application/pdf", data)
    assert out.read_as == "image" and llm.files[0][0] == "application/pdf"


@pytest.mark.parametrize(
    ("mime", "data"),
    [("image/png", b""), ("application/zip", b"PK"), ("image/png", b"x" * (8 * 1024 * 1024 + 1))],
)
def test_files_we_cannot_read_are_refused(mime, data):
    with pytest.raises(DocumentError):
        read_document(FileLLM(READING), mime, data)


def test_an_unreadable_reply_is_a_clear_error():
    with pytest.raises(DocumentError):
        read_document(FileLLM("not json"), "image/png", b"x")


def test_reading_json_round_trip():
    out = read_document(FileLLM(READING), "image/png", b"x")
    assert json.loads(json.dumps(out.as_dict()))["kind"] == "contract"
